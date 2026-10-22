"""
Laporan Hasil Audit dari Daftar Temuan — Word (dapat diedit) dan PDF.

Struktur: sampul, ringkasan eksekutif, ruang lingkup, ringkasan temuan
(tabel risiko × status + daftar), rincian per temuan (kondisi, kriteria,
rekomendasi, tanggapan, PIC, target), catatan.

Seluruh isi berasal dari data temuan; satu-satunya bagian yang (opsional)
disusun AI adalah ringkasan eksekutif, dan teks itu diedit auditor dulu.
"""

from __future__ import annotations

import re
from datetime import date
from typing import Any

from app.services.findings import FOLLOW_UP_STATUSES, RISK_LEVELS
from app.services.findings_exporter import SOURCE_LABELS, finding_code

RISK_ORDER = {r: i for i, r in enumerate(RISK_LEVELS)}
DISCLAIMER = (
    "Kondisi dan rekomendasi pada laporan ini sebagian disusun dengan bantuan analisis AI "
    "dan telah ditinjau auditor. Rujukan regulasi berasal dari dokumen sumber; verifikasi "
    "tetap diperlukan sebelum dijadikan dasar keputusan."
)
_REG_NAME_RE = re.compile(r"^(.*?)\s+(?:Penjelasan\s+)?(?:Pasal|Lampiran|BAB)\b", re.IGNORECASE)

MONTHS_ID = [
    "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember",
]


def format_date_id(value: str) -> str:
    """'2026-10-22' → '22 Oktober 2026' (teks lain dikembalikan apa adanya)."""
    try:
        d = date.fromisoformat(value)
    except (TypeError, ValueError):
        return value or ""
    return f"{d.day} {MONTHS_ID[d.month - 1]} {d.year}"


# =============================================================================
# DATA LAPORAN
# =============================================================================

def _regulation_names(kriteria: str) -> list[str]:
    names = []
    for part in re.split(r"[;\n]", kriteria or ""):
        m = _REG_NAME_RE.match(part.strip())
        if m and m.group(1).strip():
            names.append(m.group(1).strip(" —-"))
    return names


def prepare_report(findings: list[dict], options: dict) -> dict:
    """
    Pilih & urutkan temuan, hitung ringkasan, dan susun ruang lingkup.
    options: judul, auditee, disusun_oleh, tanggal (YYYY-MM-DD), ringkasan,
    statuses (list), risks (list).
    """
    # Tidak dikirim = semua; daftar kosong = tidak ada yang dipilih.
    statuses = set(FOLLOW_UP_STATUSES if options.get("statuses") is None else options["statuses"])
    risks = set(RISK_LEVELS if options.get("risks") is None else options["risks"])
    selected = [f for f in findings if f.get("status") in statuses and f.get("risiko") in risks]
    selected.sort(key=lambda f: (RISK_ORDER.get(f.get("risiko"), 9), f.get("nomor", 0)))

    matrix = {r: {s: 0 for s in FOLLOW_UP_STATUSES} for r in RISK_LEVELS}
    for f in selected:
        if f.get("risiko") in matrix and f.get("status") in matrix[f["risiko"]]:
            matrix[f["risiko"]][f["status"]] += 1

    documents: dict[str, None] = {}
    regulations: dict[str, None] = {}
    for f in selected:
        src = f.get("source") or {}
        if src.get("nama"):
            documents[f"{src['nama']} ({SOURCE_LABELS.get(src.get('type'), 'Manual')})"] = None
        for name in _regulation_names(f.get("kriteria", "")):
            regulations[name] = None

    return {
        "judul": (options.get("judul") or "").strip() or "Laporan Hasil Audit Teknologi Informasi",
        "auditee": (options.get("auditee") or "").strip(),
        "disusun_oleh": (options.get("disusun_oleh") or "").strip(),
        "tanggal": format_date_id(options.get("tanggal") or date.today().isoformat()),
        "ringkasan": (options.get("ringkasan") or "").strip(),
        "findings": selected,
        "matrix": matrix,
        "totals": {
            "total": len(selected),
            "by_risk": {r: sum(matrix[r].values()) for r in RISK_LEVELS},
            "by_status": {s: sum(matrix[r][s] for r in RISK_LEVELS) for s in FOLLOW_UP_STATUSES},
        },
        "scope": {"documents": list(documents), "regulations": list(regulations)},
    }


def _source_text(f: dict) -> str:
    src = f.get("source") or {}
    label = SOURCE_LABELS.get(src.get("type", ""), "Manual")
    detail = " · ".join(p for p in (src.get("nama"), src.get("ref")) if p)
    return f"{label}: {detail}" if detail else label


def _paragraphs(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"\n\s*\n|\n", text or "") if p.strip()]


# =============================================================================
# WORD
# =============================================================================

def build_docx(report: dict, output_path: str) -> str:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    primary = RGBColor(0x1E, 0x27, 0x61)
    muted = RGBColor(0x64, 0x74, 0x8B)
    risk_fill = {"Tinggi": "FDE2DD", "Sedang": "FFF3D6", "Rendah": "E3F3EA"}

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for level in (1, 2):
        doc.styles[f"Heading {level}"].font.color.rgb = primary

    def shade(cell, hex_fill):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tc_pr.append(shd)

    def header_row(table, labels):
        for cell, label in zip(table.rows[0].cells, labels):
            cell.text = ""
            run = cell.paragraphs[0].add_run(label)
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shade(cell, "1E2761")

    # ---------- Sampul ----------
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("LAPORAN HASIL AUDIT")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = muted
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(report["judul"])
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = primary
    for label, value in (("Auditee", report["auditee"]), ("Tanggal laporan", report["tanggal"]),
                         ("Disusun oleh", report["disusun_oleh"])):
        if value:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r1 = p.add_run(f"{label}: ")
            r1.font.color.rgb = muted
            p.add_run(value).bold = True
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------- 1. Ringkasan eksekutif ----------
    doc.add_heading("1. Ringkasan Eksekutif", level=1)
    if report["ringkasan"]:
        for para in _paragraphs(report["ringkasan"]):
            doc.add_paragraph(para)
    else:
        doc.add_paragraph("(Ringkasan eksekutif belum diisi.)").runs[0].italic = True

    # ---------- 2. Ruang lingkup ----------
    doc.add_heading("2. Ruang Lingkup", level=1)
    totals = report["totals"]
    doc.add_paragraph(
        f"Laporan ini memuat {totals['total']} temuan yang dihimpun dari hasil analisis berikut."
    )
    if report["scope"]["documents"]:
        doc.add_paragraph("Dokumen yang diuji:").runs[0].bold = True
        for d in report["scope"]["documents"]:
            doc.add_paragraph(d, style="List Bullet")
    if report["scope"]["regulations"]:
        doc.add_paragraph("Regulasi acuan:").runs[0].bold = True
        for r in report["scope"]["regulations"]:
            doc.add_paragraph(r, style="List Bullet")

    # ---------- 3. Ringkasan temuan ----------
    doc.add_heading("3. Ringkasan Temuan", level=1)
    table = doc.add_table(rows=1, cols=len(FOLLOW_UP_STATUSES) + 2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_row(table, ["Risiko", *FOLLOW_UP_STATUSES, "Jumlah"])
    for risk in RISK_LEVELS:
        cells = table.add_row().cells
        cells[0].text = risk
        shade(cells[0], risk_fill[risk])
        for i, status in enumerate(FOLLOW_UP_STATUSES, start=1):
            cells[i].text = str(report["matrix"][risk][status])
        cells[-1].text = str(totals["by_risk"][risk])
    cells = table.add_row().cells
    cells[0].text = "Jumlah"
    for i, status in enumerate(FOLLOW_UP_STATUSES, start=1):
        cells[i].text = str(totals["by_status"][status])
    cells[-1].text = str(totals["total"])
    for cell in cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True

    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    header_row(table, ["No", "Temuan", "Risiko", "Status", "PIC", "Target"])
    for f in report["findings"]:
        cells = table.add_row().cells
        values = [finding_code(f["nomor"]), f.get("judul", ""), f.get("risiko", ""),
                  f.get("status", ""), f.get("pic", "") or "-", format_date_id(f.get("target_date", "")) or "-"]
        for cell, value in zip(cells, values):
            cell.text = value
        shade(cells[2], risk_fill.get(f.get("risiko"), "FFFFFF"))

    # ---------- 4. Rincian temuan ----------
    doc.add_heading("4. Rincian Temuan", level=1)
    for f in report["findings"]:
        doc.add_heading(f"{finding_code(f['nomor'])} — {f.get('judul', '')}", level=2)
        meta = doc.add_paragraph()
        bits = [("Risiko", f.get("risiko", "")), ("Status", f.get("status", "")),
                ("PIC", f.get("pic", "") or "-"), ("Target", format_date_id(f.get("target_date", "")) or "-")]
        for i, (label, value) in enumerate(bits):
            meta.add_run(f"{label}: ").font.color.rgb = muted
            meta.add_run(value).bold = True
            if i < len(bits) - 1:
                meta.add_run("   ·   ").font.color.rgb = muted
        for label, key in (("Kondisi", "kondisi"), ("Kriteria", "kriteria"),
                           ("Rekomendasi", "rekomendasi"), ("Tanggapan / Tindak Lanjut", "tanggapan")):
            p = doc.add_paragraph()
            p.add_run(label).bold = True
            paragraphs = _paragraphs(f.get(key, "")) or ["-"]
            for para in paragraphs:
                doc.add_paragraph(para)
        src = doc.add_paragraph()
        run = src.add_run(f"Sumber: {_source_text(f)}")
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = muted

    # ---------- 5. Catatan ----------
    doc.add_heading("5. Catatan", level=1)
    doc.add_paragraph(DISCLAIMER)

    doc.save(output_path)
    return output_path


# =============================================================================
# PDF
# =============================================================================

def build_pdf(report: dict, output_path: str) -> str:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    primary = colors.HexColor("#1E2761")
    muted = colors.HexColor("#64748B")
    border = colors.HexColor("#D9E0F0")
    risk_fill = {"Tinggi": "#FDE2DD", "Sedang": "#FFF3D6", "Rendah": "#E3F3EA"}

    base = getSampleStyleSheet()
    st = {
        "kicker": ParagraphStyle("kicker", parent=base["Normal"], fontSize=12, textColor=muted,
                                 alignment=TA_CENTER, fontName="Helvetica-Bold", spaceAfter=8),
        "cover": ParagraphStyle("cover", parent=base["Title"], fontSize=22, leading=27,
                                textColor=primary, alignment=TA_CENTER, spaceAfter=18),
        "cover_meta": ParagraphStyle("cover_meta", parent=base["Normal"], fontSize=10.5, leading=15,
                                     alignment=TA_CENTER),
        "h1": ParagraphStyle("h1", parent=base["Heading1"], fontSize=14, textColor=primary, spaceBefore=10),
        "h2": ParagraphStyle("h2", parent=base["Heading2"], fontSize=11.5, textColor=primary, spaceBefore=8),
        "body": ParagraphStyle("body", parent=base["Normal"], fontSize=9.5, leading=13.5, spaceAfter=4),
        "label": ParagraphStyle("label", parent=base["Normal"], fontSize=9.5, leading=13,
                                fontName="Helvetica-Bold", spaceBefore=4),
        "small": ParagraphStyle("small", parent=base["Normal"], fontSize=8, leading=10, textColor=muted),
        "cell": ParagraphStyle("cell", parent=base["Normal"], fontSize=8.5, leading=11),
        "cell_head": ParagraphStyle("cell_head", parent=base["Normal"], fontSize=8.5, leading=11,
                                    textColor=colors.white, fontName="Helvetica-Bold"),
    }

    def esc(text: str) -> str:
        return (text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def paras(text: str, style="body"):
        return [Paragraph(esc(p), st[style]) for p in (_paragraphs(text) or ["-"])]

    def table(rows, widths, extra_styles=()):
        data = [[Paragraph(esc(str(c)), st["cell_head" if i == 0 else "cell"]) for c in row]
                for i, row in enumerate(rows)]
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), primary),
            ("GRID", (0, 0), (-1, -1), 0.5, border),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            *extra_styles,
        ]))
        return t

    story: list[Any] = [Spacer(1, 6 * cm), Paragraph("LAPORAN HASIL AUDIT", st["kicker"]),
                        Paragraph(esc(report["judul"]), st["cover"])]
    for label, value in (("Auditee", report["auditee"]), ("Tanggal laporan", report["tanggal"]),
                         ("Disusun oleh", report["disusun_oleh"])):
        if value:
            story.append(Paragraph(f"<font color='#64748B'>{label}:</font> <b>{esc(value)}</b>", st["cover_meta"]))
    story.append(PageBreak())

    story.append(Paragraph("1. Ringkasan Eksekutif", st["h1"]))
    story += paras(report["ringkasan"] or "(Ringkasan eksekutif belum diisi.)")

    totals = report["totals"]
    story.append(Paragraph("2. Ruang Lingkup", st["h1"]))
    story.append(Paragraph(
        f"Laporan ini memuat {totals['total']} temuan yang dihimpun dari hasil analisis berikut.", st["body"]))
    for title, items in (("Dokumen yang diuji", report["scope"]["documents"]),
                         ("Regulasi acuan", report["scope"]["regulations"])):
        if items:
            story.append(Paragraph(title, st["label"]))
            story += [Paragraph(f"• {esc(i)}", st["body"]) for i in items]

    story.append(Paragraph("3. Ringkasan Temuan", st["h1"]))
    rows = [["Risiko", *FOLLOW_UP_STATUSES, "Jumlah"]]
    for risk in RISK_LEVELS:
        rows.append([risk, *[report["matrix"][risk][s] for s in FOLLOW_UP_STATUSES], totals["by_risk"][risk]])
    rows.append(["Jumlah", *[totals["by_status"][s] for s in FOLLOW_UP_STATUSES], totals["total"]])
    fills = [("BACKGROUND", (0, i), (0, i), colors.HexColor(risk_fill[r])) for i, r in enumerate(RISK_LEVELS, 1)]
    story.append(table(rows, [3.8 * cm] + [3.3 * cm] * len(FOLLOW_UP_STATUSES) + [3.3 * cm], fills))
    story.append(Spacer(1, 10))

    rows = [["No", "Temuan", "Risiko", "Status", "PIC", "Target"]]
    fills = []
    for i, f in enumerate(report["findings"], start=1):
        rows.append([finding_code(f["nomor"]), f.get("judul", ""), f.get("risiko", ""), f.get("status", ""),
                     f.get("pic", "") or "-", format_date_id(f.get("target_date", "")) or "-"])
        fills.append(("BACKGROUND", (2, i), (2, i), colors.HexColor(risk_fill.get(f.get("risiko"), "#FFFFFF"))))
    story.append(table(rows, [1.6 * cm, 7.4 * cm, 1.8 * cm, 2.3 * cm, 2.6 * cm, 2.3 * cm], fills))

    story.append(Paragraph("4. Rincian Temuan", st["h1"]))
    for f in report["findings"]:
        block = [
            Paragraph(f"{finding_code(f['nomor'])} — {esc(f.get('judul', ''))}", st["h2"]),
            Paragraph(
                f"Risiko: <b>{esc(f.get('risiko', ''))}</b> · Status: <b>{esc(f.get('status', ''))}</b> · "
                f"PIC: <b>{esc(f.get('pic', '') or '-')}</b> · "
                f"Target: <b>{esc(format_date_id(f.get('target_date', '')) or '-')}</b>",
                st["small"],
            ),
        ]
        for label, key in (("Kondisi", "kondisi"), ("Kriteria", "kriteria"),
                           ("Rekomendasi", "rekomendasi"), ("Tanggapan / Tindak Lanjut", "tanggapan")):
            block.append(Paragraph(label, st["label"]))
            block += paras(f.get(key, ""))
        block.append(Paragraph(f"Sumber: {esc(_source_text(f))}", st["small"]))
        # Satu temuan tidak terpotong antarhalaman (kecuali memang lebih dari satu halaman).
        story.append(KeepTogether(block))

    story.append(Paragraph("5. Catatan", st["h1"]))
    story.append(Paragraph(esc(DISCLAIMER), st["body"]))

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(muted)
        canvas.drawString(2 * cm, 1.2 * cm, report["judul"][:90])
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Halaman {doc_.page}")
        canvas.restoreState()

    SimpleDocTemplate(
        output_path, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm, title=report["judul"],
    ).build(story, onLaterPages=footer)
    return output_path


# =============================================================================
# RINGKASAN EKSEKUTIF (AI)
# =============================================================================

SUMMARY_SYSTEM_PROMPT = """Anda adalah auditor TI senior yang menulis ringkasan eksekutif
laporan hasil audit untuk direksi bank, dalam Bahasa Indonesia formal.

ATURAN:
- Gunakan HANYA temuan yang diberikan. Jangan menambah temuan, angka, nama, atau
  pasal yang tidak ada di daftar.
- Sebut temuan dengan kodenya (mis. T-003) saat merujuk temuan tertentu.
- 3 paragraf, tanpa judul, tanpa poin-poin, tanpa markdown:
  1) kesimpulan umum dan jumlah temuan per tingkat risiko;
  2) temuan risiko tinggi paling penting dan dampaknya;
  3) prioritas tindak lanjut yang direkomendasikan.
- Paling banyak 220 kata."""


def build_summary_prompt(report: dict) -> str:
    totals = report["totals"]
    lines = [
        f"Judul audit: {report['judul']}",
        f"Auditee: {report['auditee'] or '-'}",
        f"Jumlah temuan: {totals['total']} (Tinggi {totals['by_risk']['Tinggi']}, "
        f"Sedang {totals['by_risk']['Sedang']}, Rendah {totals['by_risk']['Rendah']})",
        "",
        "DAFTAR TEMUAN:",
    ]
    for f in report["findings"]:
        lines += [
            f"- {finding_code(f['nomor'])} [{f.get('risiko')} · {f.get('status')}] {f.get('judul', '')}",
            f"  Kondisi: {(f.get('kondisi') or '-')[:500]}",
            f"  Kriteria: {(f.get('kriteria') or '-')[:200]}",
            f"  Rekomendasi: {(f.get('rekomendasi') or '-')[:300]}",
        ]
    return "\n".join(lines)


def generate_summary(report: dict, provider) -> dict:
    """Ringkasan eksekutif dari temuan; kode temuan yang tidak dikenal dicatat."""
    text = provider.generate_text(SUMMARY_SYSTEM_PROMPT, build_summary_prompt(report)).strip()
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    known = {finding_code(f["nomor"]) for f in report["findings"]}
    unknown = sorted({c for c in re.findall(r"\bT-\d{3}\b", text) if c not in known})
    notes = [f"Ringkasan menyebut kode yang tidak ada di laporan: {', '.join(unknown)} — periksa."] if unknown else []
    return {"ringkasan": text, "notes": notes}
