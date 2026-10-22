"""
Excel Exporter untuk Daftar Temuan.

Sheet "Daftar Temuan": satu baris per temuan (kondisi, kriteria, rekomendasi,
risiko, status tindak lanjut, PIC, target, tanggapan, sumber) dengan dropdown
untuk Risiko & Status. Sheet "Ringkasan": jumlah per risiko dan status.
"""

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.findings import FOLLOW_UP_STATUSES, RISK_LEVELS

HEADER_FILL = PatternFill(start_color="1E2761", end_color="1E2761", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
RISK_FILLS = {
    "Tinggi": PatternFill(start_color="FDE2DD", end_color="FDE2DD", fill_type="solid"),
    "Sedang": PatternFill(start_color="FFF3D6", end_color="FFF3D6", fill_type="solid"),
    "Rendah": PatternFill(start_color="E3F3EA", end_color="E3F3EA", fill_type="solid"),
}

COLUMNS: list[tuple[str, int]] = [
    ("No", 8),
    ("Judul Temuan", 34),
    ("Kondisi", 55),
    ("Kriteria (Regulasi)", 40),
    ("Rekomendasi", 50),
    ("Risiko", 11),
    ("Status", 14),
    ("PIC", 18),
    ("Target", 12),
    ("Tanggapan / Tindak Lanjut", 40),
    ("Sumber", 34),
]
_COL_RISK = 6
_COL_STATUS = 7

SOURCE_LABELS = {
    "compliance_review": "Compliance Review",
    "gap_analysis": "Gap Analysis",
    "manual": "Manual",
}


def finding_code(number: int) -> str:
    return f"T-{int(number):03d}"


def _source_text(source: dict) -> str:
    label = SOURCE_LABELS.get(source.get("type", ""), "Manual")
    detail = " · ".join(p for p in (source.get("nama"), source.get("ref")) if p)
    return f"{label}: {detail}" if detail else label


def _list_validation(options: tuple[str, ...], col: int, last_row: int) -> DataValidation:
    dv = DataValidation(type="list", formula1=f'"{",".join(options)}"', allow_blank=True)
    dv.error = "Pilih salah satu nilai dari daftar."
    dv.errorTitle = "Nilai tidak valid"
    dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{last_row}")
    return dv


def export_findings_to_excel(findings: list[dict[str, Any]], output_path: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Daftar Temuan"

    for col_idx, (header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 28

    for row_idx, f in enumerate(findings, start=2):
        values = [
            finding_code(f.get("nomor", row_idx - 1)),
            f.get("judul", ""),
            f.get("kondisi", ""),
            f.get("kriteria", ""),
            f.get("rekomendasi", ""),
            f.get("risiko", ""),
            f.get("status", ""),
            f.get("pic", ""),
            f.get("target_date", ""),
            f.get("tanggapan", ""),
            _source_text(f.get("source") or {}),
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            centered = col_idx in (1, _COL_RISK, _COL_STATUS, 9)
            cell.alignment = Alignment(
                horizontal="center" if centered else "left", vertical="top", wrap_text=True,
            )
        risk_fill = RISK_FILLS.get(f.get("risiko", ""))
        if risk_fill:
            ws.cell(row=row_idx, column=_COL_RISK).fill = risk_fill

    last_row = max(2, len(findings) + 1)
    ws.add_data_validation(_list_validation(RISK_LEVELS, _COL_RISK, last_row))
    ws.add_data_validation(_list_validation(FOLLOW_UP_STATUSES, _COL_STATUS, last_row))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"

    # ---------- Ringkasan ----------
    summary = wb.create_sheet("Ringkasan")
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 10
    rows: list[tuple[str, Any]] = [("Total temuan", len(findings)), ("", "")]
    rows.append(("Per tingkat risiko", ""))
    rows += [(f"  {r}", sum(1 for f in findings if f.get("risiko") == r)) for r in RISK_LEVELS]
    rows.append(("", ""))
    rows.append(("Per status tindak lanjut", ""))
    rows += [(f"  {s}", sum(1 for f in findings if f.get("status") == s)) for s in FOLLOW_UP_STATUSES]
    for i, (label, value) in enumerate(rows, start=1):
        summary.cell(row=i, column=1, value=label)
        summary.cell(row=i, column=2, value=value)
        if label and not label.startswith("  "):
            summary.cell(row=i, column=1).font = Font(bold=True)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
