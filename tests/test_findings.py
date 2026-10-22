"""
Daftar Temuan: kirim dari Compliance Review / Gap Analysis, edit, hapus, ekspor.
Fixture `client` menjalankan app di cwd sementara → data/ & outputs/ terisolasi.
"""

import io
import json
from pathlib import Path

from openpyxl import load_workbook

from app.services.findings import split_recommendation

COMP_ID = "c" * 32
GAP_ID = "d" * 32


def _seed_compliance() -> None:
    path = Path("outputs/compliance") / f"{COMP_ID}.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({
        "result_id": COMP_ID,
        "contract_filename": "MoU Vendor TI.pdf",
        "items": [
            {"klausul_ref": "Pasal 1", "klausul_text": "Ruang lingkup layanan.", "regulasi_ref": "POJK 11 Pasal 29",
             "status": "PATUH", "analisis": "Sudah sesuai."},
            {"klausul_ref": "Pasal 5", "klausul_text": "Hak audit.", "regulasi_ref": "POJK 11 Pasal 30 ayat (3)",
             "status": "TIDAK PATUH",
             "analisis": "Klausul tidak memuat hak audit OJK. REKOMENDASI: Tambahkan hak audit OJK."},
            {"klausul_ref": "Pasal 7", "klausul_text": "Kerahasiaan data.", "regulasi_ref": "POJK 11 Pasal 16",
             "status": "PATUH DENGAN CATATAN", "analisis": "Belum mengatur enkripsi."},
            {"klausul_ref": "Pasal 9", "klausul_text": "Lokasi pusat data.", "regulasi_ref": "",
             "status": "PERLU VERIFIKASI", "analisis": "Rujukan tidak dapat dicocokkan."},
            {"klausul_ref": "Pasal 10", "klausul_text": "Penutup.", "regulasi_ref": "",
             "status": "TIDAK RELEVAN", "analisis": "-"},
        ],
    }), encoding="utf-8")


def _seed_gap() -> None:
    path = Path("outputs/requirements") / f"{GAP_ID}.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({
        "result_id": GAP_ID,
        "filename": "PADK 43.pdf",
        "requirements": [
            {"persyaratan": "BPR wajib memiliki DRP.", "referensi": "Pasal 20", "catatan": "", "status_gap": "Sesuai"},
            {"persyaratan": "BPR wajib menguji DRP tahunan.", "referensi": "Pasal 21",
             "catatan": "Uji terakhir 2023.", "status_gap": "Tidak Sesuai"},
            {"persyaratan": "Kebijakan akses ditinjau berkala.", "referensi": "Pasal 22", "catatan": "",
             "status_gap": "Sebagian Sesuai"},
        ],
    }), encoding="utf-8")


def _send(client, **payload):
    return client.post("/api/findings/from-source", json=payload)


def test_kirim_semua_temuan_compliance(client):
    _seed_compliance()
    res = _send(client, source_type="compliance_review", result_id=COMP_ID)
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["candidates"] == 3 and data["skipped"] == 0
    by_ref = {f["source"]["ref"]: f for f in data["added"]}
    assert set(by_ref) == {"Pasal 5", "Pasal 7", "Pasal 9"}       # Patuh & Tidak Relevan tidak ikut

    tidak_patuh = by_ref["Pasal 5"]
    assert tidak_patuh["risiko"] == "Tinggi"
    assert tidak_patuh["status"] == "Terbuka"
    assert tidak_patuh["kondisi"] == "Klausul tidak memuat hak audit OJK."
    assert tidak_patuh["rekomendasi"] == "Tambahkan hak audit OJK."
    assert tidak_patuh["kriteria"] == "POJK 11 Pasal 30 ayat (3)"
    assert tidak_patuh["source"]["nama"] == "MoU Vendor TI.pdf"
    assert by_ref["Pasal 7"]["risiko"] == "Sedang"
    assert [f["nomor"] for f in data["added"]] == [1, 2, 3]

    # Dikirim ulang → tidak dobel.
    again = _send(client, source_type="compliance_review", result_id=COMP_ID).json()
    assert again["added"] == [] and again["skipped"] == 3
    assert client.get("/api/findings").json()["summary"]["total"] == 3


def test_kirim_baris_terpilih_gap(client):
    _seed_gap()
    res = _send(client, source_type="gap_analysis", result_id=GAP_ID, item_indices=[0, 1, 99])
    assert res.status_code == 200, res.text
    added = res.json()["added"]
    assert len(added) == 2                                        # indeks 99 diabaikan
    assert {f["risiko"] for f in added} == {"Rendah", "Tinggi"}    # Sesuai dipilih manual → Rendah
    tidak = next(f for f in added if f["risiko"] == "Tinggi")
    assert "Uji terakhir 2023." in tidak["kondisi"]
    assert tidak["kriteria"].startswith("Pasal 21")


def test_kirim_tanpa_temuan_atau_hasil_tidak_ada(client):
    _seed_gap()
    assert _send(client, source_type="gap_analysis", result_id=GAP_ID, item_indices=[99]).status_code == 422
    assert _send(client, source_type="compliance_review", result_id="e" * 32).status_code == 404
    assert _send(client, source_type="lainnya", result_id=GAP_ID).status_code == 422


def test_ubah_temuan_dan_validasi(client):
    _seed_compliance()
    finding = _send(client, source_type="compliance_review", result_id=COMP_ID).json()["added"][0]
    res = client.patch(f"/api/findings/{finding['id']}", json={
        "status": "Dalam Proses", "pic": "Divisi TI", "target_date": "2026-12-31", "risiko": "Rendah",
    })
    assert res.status_code == 200, res.text
    updated = res.json()
    assert (updated["status"], updated["pic"], updated["target_date"], updated["risiko"]) == \
        ("Dalam Proses", "Divisi TI", "2026-12-31", "Rendah")
    assert updated["kondisi"] == finding["kondisi"]               # field lain tidak berubah

    assert client.patch(f"/api/findings/{finding['id']}", json={"risiko": "Kritis"}).status_code == 422
    assert client.patch(f"/api/findings/{finding['id']}", json={"target_date": "31/12/2026"}).status_code == 422
    assert client.patch(f"/api/findings/{'0' * 32}", json={"pic": "x"}).status_code == 404

    summary = client.get("/api/findings").json()["summary"]
    assert summary["by_status"]["Dalam Proses"] == 1


def test_temuan_manual_dan_nomor_tidak_dipakai_ulang(client):
    assert client.post("/api/findings", json={"judul": "  "}).status_code == 422
    first = client.post("/api/findings", json={"judul": "Log akses tidak ditinjau"}).json()
    assert first["nomor"] == 1 and first["source"]["type"] == "manual"
    assert client.delete(f"/api/findings/{first['id']}").status_code == 200
    assert client.delete(f"/api/findings/{first['id']}").status_code == 404
    second = client.post("/api/findings", json={"judul": "Temuan kedua"}).json()
    assert second["nomor"] == 2                                   # T-001 tidak dipakai ulang


def test_ekspor_excel(client):
    assert client.get("/api/findings/export").status_code == 422   # masih kosong
    _seed_compliance()
    _send(client, source_type="compliance_review", result_id=COMP_ID)
    res = client.get("/api/findings/export")
    assert res.status_code == 200, res.text
    assert "daftar_temuan_" in res.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(res.content))
    ws = wb["Daftar Temuan"]
    headers = [c.value for c in ws[1]]
    assert headers[:7] == ["No", "Judul Temuan", "Kondisi", "Kriteria (Regulasi)", "Rekomendasi", "Risiko", "Status"]
    assert ws["A2"].value == "T-001"
    assert ws.max_row == 4
    assert {ws.cell(row=r, column=6).value for r in range(2, 5)} == {"Tinggi", "Sedang"}
    ringkasan = {row[0].value: row[1].value for row in wb["Ringkasan"].iter_rows()}
    assert ringkasan["Total temuan"] == 3


def test_split_rekomendasi():
    assert split_recommendation("A belum ada. REKOMENDASI: Tambahkan A.") == ("A belum ada.", "Tambahkan A.")
    assert split_recommendation("Tanpa saran.") == ("Tanpa saran.", "")
    assert split_recommendation("") == ("", "")
