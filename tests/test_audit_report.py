"""
Laporan Hasil Audit dari Daftar Temuan: susunan data, Word, PDF, pengaturan
tersimpan, dan draf ringkasan eksekutif (AI tiruan).
"""

import io
from unittest.mock import patch

import fitz  # PyMuPDF
from docx import Document

from app.services.ai_service import AIProvider
from app.services.audit_report import format_date_id, prepare_report


class _SummaryProvider(AIProvider):
    def __init__(self, text):
        self.text = text
        self.prompts = []

    def generate_text(self, system_prompt, user_prompt, expect_json=False):
        self.prompts.append(user_prompt)
        return self.text


def _finding(nomor, risiko, status, **extra):
    return {
        "nomor": nomor, "judul": f"Temuan {nomor}", "risiko": risiko, "status": status,
        "kondisi": f"Kondisi {nomor}", "kriteria": "POJK 11 03/2022 Pasal 30 ayat (3)",
        "rekomendasi": f"Rekomendasi {nomor}", "pic": "", "target_date": "", "tanggapan": "",
        "source": {"type": "compliance_review", "nama": "MoU.pdf", "ref": f"Pasal {nomor}"},
        **extra,
    }


def _seed(client):
    for judul, risiko, status in (
        ("Subkontrak tanpa persetujuan", "Tinggi", "Terbuka"),
        ("RPB belum diuji", "Sedang", "Dalam Proses"),
        ("Log akses tidak ditinjau", "Rendah", "Selesai"),
    ):
        res = client.post("/api/findings", json={
            "judul": judul, "risiko": risiko, "status": status,
            "kondisi": f"Kondisi: {judul}.", "kriteria": "POJK 11 03/2022 Pasal 30 ayat (3)",
            "rekomendasi": f"Perbaiki: {judul}.", "pic": "Divisi TI", "target_date": "2026-12-31",
        })
        assert res.status_code == 201, res.text


OPTIONS = {
    "judul": "Audit Penyedia Jasa TI 2026", "auditee": "Divisi TI",
    "disusun_oleh": "SKAI", "tanggal": "2026-10-22", "ringkasan": "Ringkasan uji.",
}


def test_prepare_report_filter_urutan_ringkasan_dan_lingkup():
    findings = [
        _finding(1, "Rendah", "Terbuka"),
        _finding(2, "Tinggi", "Selesai"),
        _finding(3, "Tinggi", "Terbuka", kriteria="PADK 43 PADK03 2025 Lampiran I / BAB V; POJK 11 03/2022 Pasal 35"),
        _finding(4, "Sedang", "Dalam Proses", source={"type": "gap_analysis", "nama": "PADK 43.pdf"}),
    ]
    report = prepare_report(findings, {"statuses": ["Terbuka", "Dalam Proses"], "tanggal": "2026-10-22"})
    assert [f["nomor"] for f in report["findings"]] == [3, 4, 1]   # risiko tertinggi dulu; Selesai dibuang
    assert report["matrix"]["Tinggi"] == {"Terbuka": 1, "Dalam Proses": 0, "Selesai": 0}
    assert report["totals"]["total"] == 3 and report["totals"]["by_risk"]["Sedang"] == 1
    assert report["scope"]["documents"] == ["MoU.pdf (Compliance Review)", "PADK 43.pdf (Gap Analysis)"]
    # Urutan mengikuti temuan (risiko tertinggi dulu), tanpa duplikat.
    assert report["scope"]["regulations"] == ["PADK 43 PADK03 2025", "POJK 11 03/2022"]
    assert report["tanggal"] == "22 Oktober 2026"
    assert report["judul"] == "Laporan Hasil Audit Teknologi Informasi"   # default bila kosong


def test_format_tanggal_indonesia():
    assert format_date_id("2027-03-01") == "1 Maret 2027"
    assert format_date_id("") == ""
    assert format_date_id("bukan tanggal") == "bukan tanggal"


def test_unduh_word_berisi_semua_bagian(client):
    _seed(client)
    res = client.post("/api/findings/report?format=docx", json=OPTIONS)
    assert res.status_code == 200, res.text
    assert res.headers["content-disposition"].endswith('.docx"')
    doc = Document(io.BytesIO(res.content))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings[:4] == ["1. Ringkasan Eksekutif", "2. Ruang Lingkup", "3. Ringkasan Temuan", "4. Rincian Temuan"]
    assert headings[4].startswith("T-001 — Subkontrak")                    # Tinggi paling atas
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Audit Penyedia Jasa TI 2026" in text and "Ringkasan uji." in text
    assert "31 Desember 2026" in text
    assert len(doc.tables) == 2


def test_unduh_pdf_dan_filter(client):
    _seed(client)
    res = client.post("/api/findings/report?format=pdf", json={**OPTIONS, "statuses": ["Terbuka"]})
    assert res.status_code == 200, res.text
    assert res.content.startswith(b"%PDF")
    text = "".join(page.get_text() for page in fitz.open(stream=res.content, filetype="pdf"))
    assert "Subkontrak tanpa persetujuan" in text
    assert "RPB belum diuji" not in text                                   # Dalam Proses tidak dipilih

    none = client.post("/api/findings/report?format=pdf", json={**OPTIONS, "risks": []})
    assert none.status_code == 422
    assert client.post("/api/findings/report?format=xlsx", json=OPTIONS).status_code == 422


def test_pengaturan_laporan_diingat(client):
    assert client.get("/api/findings/report/meta").json()["judul"] == ""
    _seed(client)
    client.post("/api/findings/report?format=docx", json={**OPTIONS, "risks": ["Tinggi"]})
    meta = client.get("/api/findings/report/meta").json()
    assert meta["judul"] == "Audit Penyedia Jasa TI 2026"
    assert meta["ringkasan"] == "Ringkasan uji."
    assert meta["risks"] == ["Tinggi"]
    assert client.get("/api/findings").json()["summary"]["total"] == 3     # temuan tidak tersentuh


def test_draf_ringkasan_ai_dan_kode_tak_dikenal(client):
    _seed(client)
    provider = _SummaryProvider("Audit menemukan **3 temuan**. Prioritas: T-001 dan T-009.")
    with patch("app.api.common.get_ai_provider", return_value=provider):
        res = client.post("/api/findings/report/summary", json=OPTIONS)
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["ringkasan"] == "Audit menemukan 3 temuan. Prioritas: T-001 dan T-009."   # markdown dibuang
    assert data["notes"] and "T-009" in data["notes"][0]
    prompt = provider.prompts[0]
    assert "T-001 [Tinggi · Terbuka] Subkontrak tanpa persetujuan" in prompt
    assert "Kondisi: Subkontrak tanpa persetujuan." in prompt
